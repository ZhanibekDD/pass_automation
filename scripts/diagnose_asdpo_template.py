"""
Диагностика боевого шаблона АСДПО: только чтение, файл не сохраняется.

Запуск из корня проекта pass_automation:
  python scripts/diagnose_asdpo_template.py
"""

from __future__ import annotations

import sys
from pathlib import Path

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except AttributeError:
        pass

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent.parent
TEMPLATE_PATH = BASE_DIR / "data" / "templates" / "asdpo_template.xlsx"

SHEET_NAMES = ("Заявка", "Реестр", "СубПодрядчик")


def main() -> None:
    if not TEMPLATE_PATH.exists():
        raise SystemExit(f"Файл не найден: {TEMPLATE_PATH}")

    wb = load_workbook(
        TEMPLATE_PATH,
        read_only=True,
        data_only=True,
    )

    try:
        print("Файл:", TEMPLATE_PATH.resolve())
        print()
        print("=== Имена листов ===")
        for i, name in enumerate(wb.sheetnames, start=1):
            print(f"  {i}. {name!r}")
        print()

        for sheet_name in SHEET_NAMES:
            if sheet_name not in wb.sheetnames:
                print(f"=== Лист {sheet_name!r} — нет в книге, пропуск ===\n")
                continue

            ws = wb[sheet_name]
            print(f"=== Лист {sheet_name!r} ===")
            dim = getattr(ws, "dimensions", None)
            if dim is not None:
                print(f"dimensions (openpyxl): {dim}")
            else:
                print("dimensions: в режиме read_only недоступны (см. ячейки ниже)")
            print("Непустые ячейки в диапазоне A1:Z15 (адрес = значение):")
            for row in ws.iter_rows(
                min_row=1,
                max_row=15,
                min_col=1,
                max_col=26,
                values_only=False,
            ):
                for cell in row:
                    val = cell.value
                    if val is None:
                        continue
                    if isinstance(val, str) and not val.strip():
                        continue
                    print(f"  {cell.coordinate} = {val!r}")
            print()

        if "Реестр" in wb.sheetnames:
            ws = wb["Реестр"]
            print("=== Реестр: строка 5, колонки A:Z (значения) ===")
            rows = ws.iter_rows(
                min_row=5,
                max_row=5,
                min_col=1,
                max_col=26,
                values_only=True,
            )
            row5 = next(rows)
            for col_idx, val in enumerate(row5, start=1):
                col = get_column_letter(col_idx)
                print(f"  {col}5 = {val!r}")
            print()

    finally:
        wb.close()

    print("Готово. Файл шаблона не изменялся.")


if __name__ == "__main__":
    main()
